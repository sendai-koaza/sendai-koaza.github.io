"""Build the public place-name list pages from the maintained workbook."""

import html
import json
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\地名すべて\仙台市小字地図_小字一覧.xlsx")
ROOT = Path(__file__).resolve().parents[1]
LIST_ROOT = ROOT / "koaza-list"

MAP_ALIASES = {
    "南小泉": {"五ツ谷": "五ッ谷", "五ツ谷二番": "五ッ谷二番"},
    "蒲町": {"川崎": "川嵜", "川崎北": "川嵜北"},
    "伊在": {"土府": "𡈽府", "南土府": "南𡈽府"},
    "長喜城": {"御蔵堀": "御藏堀", "鉄炮前": "鉄砲前"},
    "南目": {"南草刈谷地": "南草苅谷地", "朝刈場": "朝苅塲", "柳町": "栁町", "草刈谷地": "草苅谷地"},
}

PAGES = [
    {"sheet": "仙台城下", "slug": "", "label": "仙台城下", "oaza": "仙台"},
    {"sheet": "南小泉", "slug": "minamikoizumi", "label": "南小泉村", "oaza": "南小泉"},
    {"sheet": "蒲町", "slug": "kabamachi", "label": "蒲町村", "oaza": "蒲町"},
    {"sheet": "伊在", "slug": "izai", "label": "伊在村", "oaza": "伊在"},
    {"sheet": "霞目", "slug": "kasuminome", "label": "霞目村", "oaza": "霞目"},
    {"sheet": "長喜城", "slug": "choukijou", "label": "長喜城村", "oaza": "長喜城"},
    {"sheet": "六丁目", "slug": "rokuchoume", "label": "六丁目村", "oaza": "六丁目"},
    {"sheet": "南目", "slug": "minamime", "label": "南目村", "oaza": "南目"},
    {"sheet": "苦竹", "slug": "nigatake", "label": "苦竹村", "oaza": "苦竹"},
]

COMMON_NOTES = [
    "地名の表記揺れ・別名・誤記などがある場合は備考に記載しています。その場合、(揺)・(別)・(誤)などを頭に付けます。",
    "ただし、丁・町の表記揺れは省略し、現在一般的な表記を採用しています。また、新・旧字体の表記揺れも省略し、特に必要がなければ新字体で表記しています。",
    "順番は陸前国宮城郡各村字調書の記載の通りですが、地図にあって調書にない地名も備考に記載のうえ適宜挿入しています。",
    "調書にあって地図にない地名はページ下部に記載しています（該当する地名が存在しない村もあります）。",
    "地名をクリックすると、地図の該当地域に飛びます。",
]


def value(cell):
    return "" if cell.value is None else str(cell.value).strip()


def find_header(sheet, headers):
    for row in sheet.iter_rows():
        for cell in row:
            if value(cell) != headers[0]:
                continue
            if all(value(sheet.cell(cell.row, cell.column + index)) == name for index, name in enumerate(headers)):
                return cell.row, cell.column
    return None


def find_sparse_header(sheet, headers):
    for row_number in range(1, sheet.max_row + 1):
        columns = {}
        for column_number in range(1, sheet.max_column + 1):
            cell_value = value(sheet.cell(row_number, column_number))
            if cell_value in headers:
                columns[cell_value] = column_number
        if len(columns) == len(headers):
            return row_number, columns
    return None


def table_rows(sheet, header_row, start_column, headers, end_row=None):
    rows = []
    for row_number in range(header_row + 1, (end_row or sheet.max_row) + 1):
        values = [value(sheet.cell(row_number, start_column + index)) for index in range(len(headers))]
        if any(values):
            rows.append(dict(zip(headers, values)))
    return rows


def sparse_table_rows(sheet, header_row, columns, headers, required_header=None):
    rows = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        row = {header: value(sheet.cell(row_number, columns[header])) for header in headers}
        if any(row.values()) and (not required_header or row[required_header]):
            rows.append(row)
    return rows


def navigation(sheet):
    groups = []
    for row_number in range(3, 9):
        ward = value(sheet.cell(row_number, 1))
        places = [item.strip() for item in value(sheet.cell(row_number, 2)).split("、") if item.strip()]
        if places:
            groups.append({"ward": ward, "places": places})
    return groups


def find_unmapped_marker(sheet, after_row):
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row > after_row and "地図に存在しない地名" in value(cell):
                return cell.row
    return None


def page_path(config):
    return "./koaza-list/" if not config["slug"] else "./koaza-list/" + config["slug"] + "/"


def page_links(config):
    links = {}
    for candidate in PAGES:
        if candidate["label"] == config["label"]:
            links[candidate["label"]] = "#list-title"
        elif not config["slug"]:
            links[candidate["label"]] = "./" + candidate["slug"] + "/"
        elif not candidate["slug"]:
            links[candidate["label"]] = "../"
        else:
            links[candidate["label"]] = "../" + candidate["slug"] + "/"
    return links


def notes_for(config):
    notes = list(COMMON_NOTES)
    if config["oaza"] == "仙台":
        notes[2] = notes[2].replace("陸前国宮城郡各村字調書", "陸前国仙臺区各村町調書")
    return notes


def page_html(config, title):
    is_root = not config["slug"]
    prefix = "../" if is_root else "../../"
    list_prefix = "./" if is_root else "../"
    data_prefix = "./"
    return f'''<!doctype html>
<html lang="ja">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}の小字一覧｜仙台市小字地図</title>
  <link rel="icon" type="image/png" sizes="512x512" href="{prefix}resources/favicon.png">
  <link rel="apple-touch-icon" href="{prefix}resources/favicon.png">
  <link rel="stylesheet" href="{prefix}resources/content-pages.css">
</head>
<body class="content-page koaza-list-page">
  <header class="content-header">
    <div class="content-header__inner">
      <a class="content-brand" href="{prefix}">仙台市小字地図</a>
      <nav class="content-nav" aria-label="サイト内メニュー"><a href="{prefix}">地図</a><a href="{list_prefix}" aria-current="page">小字一覧</a><a href="{prefix}about/">このサイトについて</a></nav>
    </div>
  </header>
  <main class="content-main">
    <h1>小字一覧</h1>
    <section class="content-section"><h2>地域・村を選ぶ</h2><div id="area-navigation" class="list-navigation" aria-label="地域・村の一覧"></div></section>
    <section class="content-section content-card list-notes"><h2>表記について</h2><div id="list-notes"></div></section>
    <section class="content-section"><h2 id="list-title">{html.escape(title)}の小字一覧</h2><div class="list-table-wrap"><table id="main-list" class="list-table"></table></div></section>
    <section id="extra-section" class="content-section" hidden><h2>おまけ</h2><p>調書に存在し、地図に存在しない地名です。</p><div class="list-table-wrap"><table id="extra-list" class="list-table"></table></div></section>
    <section id="references-section" class="content-section"><h2>参考文献</h2><div class="list-table-wrap"><table id="references-list" class="list-table references-table"></table></div></section>
  </main>
  <footer class="content-footer"><a href="{prefix}">仙台市小字地図へ戻る</a></footer>
  <script src="{data_prefix}data.js"></script><script src="{list_prefix}list.js"></script>
</body>
</html>
'''


def build_page(sheet, config, shared_navigation):
    main_headers = ["地名", "読み", "由来", "備考"]
    reference_headers = ["番号", "資料名", "リンク"]
    main_header = find_header(sheet, main_headers)
    if not main_header:
        raise ValueError(f"Main list header not found in {sheet.title}")
    main_header_row, main_start_column = main_header
    extra_headers = ["地名", "読み", "場所", "由来", "備考"]
    extra_header = find_header(sheet, extra_headers)
    reference_header = find_sparse_header(sheet, reference_headers)
    unmapped_marker_row = find_unmapped_marker(sheet, main_header_row)
    if unmapped_marker_row:
        main_end_row = unmapped_marker_row - 1
    elif reference_header and reference_header[0] >= main_header_row + 10:
        main_end_row = reference_header[0] - 1
    else:
        main_end_row = None
    if extra_header:
        extra_rows = table_rows(sheet, *extra_header, extra_headers)
    elif unmapped_marker_row:
        extra_rows = table_rows(sheet, unmapped_marker_row, main_start_column, main_headers)
    else:
        extra_rows = []
    title = "仙台城下" if config["oaza"] == "仙台" else sheet.title
    directory = LIST_ROOT if not config["slug"] else LIST_ROOT / config["slug"]
    data = {
        "title": title,
        "navigation": shared_navigation,
        "notes": notes_for(config),
        "main": table_rows(sheet, main_header_row, main_start_column, main_headers, main_end_row),
        "extra": extra_rows,
        "references": sparse_table_rows(sheet, reference_header[0], reference_header[1], reference_headers, "資料名") if reference_header else [],
        "mapAliases": MAP_ALIASES.get(config["oaza"], {}),
        "mapBase": "../" if not config["slug"] else "../../",
        "oaza": config["oaza"],
        "pageLinks": page_links(config),
    }
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "data.js").write_text("window.KOAZA_LIST_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    (directory / "index.html").write_text(page_html(config, title), encoding="utf-8")
    return data


def main():
    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    shared_navigation = navigation(workbook["仙台城下"])
    built = {}
    for config in PAGES:
        data = build_page(workbook[config["sheet"]], config, shared_navigation)
        built[config["oaza"]] = {
            "path": page_path(config),
            "items": {item["地名"]: index for index, item in enumerate(data["main"])},
            "mapAliases": data["mapAliases"],
        }
        print(f"Updated {config['sheet']}: {len(data['main'])} main rows.")
    (LIST_ROOT / "pages.js").write_text(
        "window.KOAZA_LIST_PAGE_INDEX = " + json.dumps(built, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
